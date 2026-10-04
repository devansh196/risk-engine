"""Write the daily risk report to a formatted Excel workbook.

Change columns, weights, P&L, totals, utilisation and excess collateral are written as
Excel FORMULAS, so the workbook stays live if someone edits a number.
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "1F3A5F"
HEADER = PatternFill("solid", fgColor=NAVY)
SUBTLE = PatternFill("solid", fgColor="EEF2F7")
STATUS_FILL = {"INFO": PatternFill("solid", fgColor="E2F0D9"),
               "WARNING": PatternFill("solid", fgColor="FFF2CC"),
               "ALERT": PatternFill("solid", fgColor="F8CBAD")}
ZONE_FILL = {"GREEN": STATUS_FILL["INFO"], "YELLOW": STATUS_FILL["WARNING"], "RED": STATUS_FILL["ALERT"]}
WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
INR = '#,##0;[Red]-#,##0'
PCT = '0.0%;[Red]-0.0%'
PRICE = '#,##0.00'


def _header(ws, row, labels, col=1):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=label)
        c.fill, c.font = HEADER, WHITE_BOLD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, color="595959")


def write_report(path, t, y, asset_pnl, units, bt_summary, bt_daily, stress, history,
                 comments, var_limit, by_class=None):
    wb = Workbook()
    d_t, d_y = t["date"].date(), y["date"].date()
    sub = f"As of {d_t}  |  compared with {d_y}  |  99% confidence  |  all values in INR"

    # ---------------- Summary ----------------
    ws = wb.active
    ws.title = "Summary"
    _title(ws, "Daily Risk & Margin Report", sub)
    _header(ws, 4, ["Metric", f"Previous ({d_y})", f"Today ({d_t})", "Change", "% Change"])
    kpis = [
        ("Portfolio market value", y["total"], t["total"], INR),
        ("1-day VaR, historical", y["hist_var"], t["hist_var"], INR),
        ("1-day Expected Shortfall, historical", y["hist_es"], t["hist_es"], INR),
        ("1-day VaR, parametric", y["param_var"], t["param_var"], INR),
        ("1-day VaR, EWMA", y["ewma_var"], t["ewma_var"], INR),
        ("VaR limit", var_limit, var_limit, INR),
        ("VaR limit utilisation", None, None, PCT),             # formula
        ("Initial margin, recent calibration", y["im"]["Recent (last 2y)"], t["im"]["Recent (last 2y)"], INR),
        ("Initial margin, stressed (2020)", y["im"]["Stressed (2020)"], t["im"]["Stressed (2020)"], INR),
        ("Initial margin required (APC blend)", y["im_required"], t["im_required"], INR),
        ("Collateral after haircuts", y["collateral_after_haircut"], t["collateral_after_haircut"], INR),
        ("Excess / (shortfall) collateral", None, None, INR),   # formula
    ]
    rows = {}
    for i, (name, prev, today, fmt) in enumerate(kpis, start=5):
        rows[name] = i
        ws.cell(row=i, column=1, value=name)
        for col, v in ((2, prev), (3, today)):
            cell = ws.cell(row=i, column=col, value=None if v is None or pd.isna(v) else float(v))
            cell.number_format = fmt
        ws.cell(row=i, column=4, value=f"=C{i}-B{i}").number_format = INR if fmt == INR else PCT
        ws.cell(row=i, column=5, value=f'=IF(OR(B{i}="",B{i}=0),"",D{i}/B{i})').number_format = PCT
        for col in range(1, 6):
            ws.cell(row=i, column=col).border = THIN
    r_var, r_lim, r_util = rows["1-day VaR, historical"], rows["VaR limit"], rows["VaR limit utilisation"]
    r_im, r_col, r_exc = (rows["Initial margin required (APC blend)"], rows["Collateral after haircuts"],
                          rows["Excess / (shortfall) collateral"])
    for col in "BC":
        ws[f"{col}{r_util}"] = f"={col}{r_var}/{col}{r_lim}"
        ws[f"{col}{r_exc}"] = f"={col}{r_col}-{col}{r_im}"
    ws[f"D{r_util}"].number_format = PCT
    ws[f"E{r_util}"] = ""
    for name in ("VaR limit utilisation", "Excess / (shortfall) collateral", "Initial margin required (APC blend)"):
        ws.cell(row=rows[name], column=1).font = BOLD

    r = r_exc + 1
    pnl_row = r
    ws.cell(row=r, column=1, value="Daily P&L").font = BOLD
    ws.cell(row=r, column=3, value="=Positions!I{}".format(len(units) + 5)).number_format = INR
    r += 1
    ws.cell(row=r, column=1, value=f"VaR exceptions, last {bt_summary['Days']} days").font = BOLD
    ws.cell(row=r, column=3, value=int(bt_summary["Exceptions"]))
    ws.cell(row=r, column=4, value=bt_summary["Basel zone"]).fill = ZONE_FILL[bt_summary["Basel zone"]]

    r += 2
    ws.cell(row=r, column=1, value="Commentary").font = Font(bold=True, size=12, color=NAVY)
    r += 1
    _header(ws, r, ["Status", "Comment"])
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    for status, text in comments:
        r += 1
        s = ws.cell(row=r, column=1, value=status)
        s.fill, s.font = STATUS_FILL[status], BOLD
        s.alignment = Alignment(horizontal="center", vertical="top")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        c = ws.cell(row=r, column=2, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 15 * (len(text) // 88 + 1) + 2
    _widths(ws, [40, 20, 20, 16, 12])
    ws.freeze_panes = "A5"

    # ---------------- Positions ----------------
    ws = wb.create_sheet("Positions")
    _title(ws, "Positions, P&L and VaR contribution", sub)
    cols = ["Asset", "Class", "Units", f"Price {d_y}", f"Price {d_t}", "Market value", "Weight",
            "Daily return", "Daily P&L", "Standalone VaR", f"VaR contribution {d_y}",
            f"VaR contribution {d_t}", "Change in contribution"]
    _header(ws, 4, cols)
    from .scenarios import asset_class
    first = 5
    for i, tk in enumerate(units.index, start=first):
        ws.cell(row=i, column=1, value=tk)
        ws.cell(row=i, column=2, value=asset_class(tk))
        ws.cell(row=i, column=3, value=float(units[tk])).number_format = '#,##0.000'
        ws.cell(row=i, column=4, value=float(y["prices"][tk])).number_format = PRICE
        ws.cell(row=i, column=5, value=float(t["prices"][tk])).number_format = PRICE
        ws.cell(row=i, column=6, value=f"=C{i}*E{i}").number_format = INR
        ws.cell(row=i, column=8, value=f"=E{i}/D{i}-1").number_format = PCT
        ws.cell(row=i, column=9, value=f"=C{i}*(E{i}-D{i})").number_format = INR
        ws.cell(row=i, column=10, value=float(t["standalone_var"][tk])).number_format = INR
        ws.cell(row=i, column=11, value=float(y["component_var"][tk])).number_format = INR
        ws.cell(row=i, column=12, value=float(t["component_var"][tk])).number_format = INR
        ws.cell(row=i, column=13, value=f"=L{i}-K{i}").number_format = INR
    last = first + len(units) - 1
    tot = last + 1
    for i in range(first, last + 1):
        ws.cell(row=i, column=7, value=f"=F{i}/$F${tot}").number_format = PCT
    ws.cell(row=tot, column=1, value="TOTAL").font = BOLD
    for col in "FGIJKLM":
        c = ws[f"{col}{tot}"]
        c.value = f"=SUM({col}{first}:{col}{last})"
        c.font, c.fill = BOLD, SUBTLE
        c.number_format = PCT if col == "G" else INR
    ws[f"H{tot}"] = f"=I{tot}/SUMPRODUCT(C{first}:C{last},D{first}:D{last})"
    ws[f"H{tot}"].number_format, ws[f"H{tot}"].font, ws[f"H{tot}"].fill = PCT, BOLD, SUBTLE
    ws.cell(row=tot + 2, column=1, value="Diversification benefit (standalone VaR sum - parametric VaR)")
    ws.cell(row=tot + 2, column=10, value=f"=J{tot}-L{tot}").number_format = INR
    ws.cell(row=tot + 3, column=1, value="VaR contributions are component VaRs: they add up to the "
                                         "parametric VaR. Positive change = asset added risk today.").font = Font(italic=True, color="595959")
    _widths(ws, [16, 9, 12, 13, 13, 14, 9, 11, 13, 14, 16, 16, 14])
    ws.freeze_panes = "B5"

    # ---------------- Margin ----------------
    ws = wb.create_sheet("Margin")
    _title(ws, "Initial margin and collateral", sub)
    _header(ws, 4, ["Initial margin (10-day MPOR)", f"Previous ({d_y})", f"Today ({d_t})", "Change"])
    for i, k in enumerate(t["im"], start=5):
        ws.cell(row=i, column=1, value=k)
        for col, snap in ((2, y), (3, t)):
            v = snap["im"][k]
            ws.cell(row=i, column=col, value=None if pd.isna(v) else float(v)).number_format = INR
        ws.cell(row=i, column=4, value=f"=C{i}-B{i}").number_format = INR
    r = 5 + len(t["im"]) + 1
    if by_class is not None:
        _header(ws, r, ["IM by asset class (today)", "Standalone IM"])
        for k, v in by_class.items():
            r += 1
            ws.cell(row=r, column=1, value=k)
            ws.cell(row=r, column=2, value=float(v)).number_format = INR
        r += 2
    _header(ws, r, ["Collateral (today)", "Price driver", "Units", "Price", "Market value", "Haircut",
                    "Value after haircut"])
    c0 = r + 1
    for name, row in t["collateral"].iterrows():
        r += 1
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=row["Price driver"])
        if row["Units"] is not None and not pd.isna(row["Units"]):
            ws.cell(row=r, column=3, value=float(row["Units"])).number_format = '#,##0.000'
            ws.cell(row=r, column=4, value=float(row["Price"])).number_format = PRICE
            ws.cell(row=r, column=5, value=f"=C{r}*D{r}").number_format = INR
        else:
            ws.cell(row=r, column=5, value=float(row["Market value"])).number_format = INR
        ws.cell(row=r, column=6, value=float(row["Haircut"])).number_format = '0%'
        ws.cell(row=r, column=7, value=f"=E{r}*(1-F{r})").number_format = INR
    r += 1
    ws.cell(row=r, column=1, value="TOTAL").font = BOLD
    for col in "EG":
        ws[f"{col}{r}"] = f"=SUM({col}{c0}:{col}{r - 1})"
        ws[f"{col}{r}"].number_format, ws[f"{col}{r}"].font = INR, BOLD
    r += 2
    ws.cell(row=r, column=1, value="Initial margin required (APC blend)")
    ws.cell(row=r, column=7, value=f"=C{5 + list(t['im']).index(next(k for k in t['im'] if k.startswith('APC')))}").number_format = INR
    ws.cell(row=r + 1, column=1, value="Excess / (shortfall) collateral").font = BOLD
    ws.cell(row=r + 1, column=7, value=f"=G{r - 2}-G{r}").number_format = INR
    ws.cell(row=r + 1, column=7).font = BOLD
    _widths(ws, [38, 16, 16, 12, 14, 9, 18])

    # ---------------- Stress ----------------
    ws = wb.create_sheet("Stress")
    _title(ws, "Stress tests on today's portfolio", sub)
    _header(ws, 4, ["Scenario", "Type", "P&L", "% of portfolio", "x 1-day VaR", "Covered by collateral"])
    for i, (name, row) in enumerate(stress.iterrows(), start=5):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=row["Type"])
        ws.cell(row=i, column=3, value=float(row["P&L"])).number_format = INR
        ws.cell(row=i, column=4, value=float(row["% of portfolio"])).number_format = PCT
        ws.cell(row=i, column=5, value=f"=IF(C{i}<0,-C{i}/Summary!C{r_var},0)").number_format = '0.0"x"'
        ws.cell(row=i, column=6, value=f'=IF(C{i}<0,MIN(1,Summary!C{r_col}/-C{i}),"")').number_format = '0%'
    _widths(ws, [36, 14, 14, 14, 12, 20])
    n = len(stress)
    ch = BarChart(); ch.type = "bar"; ch.style = 10
    ch.title = "Stress scenario P&L (INR)"; ch.legend = None
    ch.add_data(Reference(ws, min_col=3, min_row=4, max_row=4 + n), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=4 + n))
    ch.height, ch.width = 9, 18
    ch.y_axis.numFmt = '#,##0'
    ws.add_chart(ch, "H4")

    # ---------------- VaR history ----------------
    ws = wb.create_sheet("VaR history")
    _title(ws, f"VaR vs actual P&L, last {len(history)} days", sub)
    _header(ws, 4, ["Date", "Actual P&L", "Loss", "Historical VaR", "Parametric VaR", "EWMA VaR", "Exception"])
    for i, (d, row) in enumerate(history.iterrows(), start=5):
        ws.cell(row=i, column=1, value=d.date()).number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=2, value=float(row["Actual P&L"])).number_format = INR
        ws.cell(row=i, column=3, value=f"=MAX(0,-B{i})").number_format = INR
        for col, key in ((4, "Historical VaR"), (5, "Parametric VaR"), (6, "EWMA VaR")):
            ws.cell(row=i, column=col, value=float(row[key])).number_format = INR
        e = ws.cell(row=i, column=7, value=f'=IF(C{i}>D{i},"YES","")')
        if row["Exception"]:
            for col in range(1, 8):
                ws.cell(row=i, column=col).fill = STATUS_FILL["ALERT"]
    _widths(ws, [12, 13, 12, 15, 15, 13, 10])
    ws.freeze_panes = "A5"
    m = 4 + len(history)
    lc = LineChart(); lc.title = "Daily loss vs 99% 1-day VaR"; lc.style = 12
    lc.add_data(Reference(ws, min_col=3, max_col=6, min_row=4, max_row=m), titles_from_data=True)
    lc.set_categories(Reference(ws, min_col=1, min_row=5, max_row=m))
    lc.height, lc.width = 9, 22
    lc.y_axis.numFmt = '#,##0'
    lc.x_axis.number_format = "dd-mmm"
    ws.add_chart(lc, "I4")

    # ---------------- Backtest ----------------
    ws = wb.create_sheet("Backtest")
    _title(ws, f"Historical VaR backtest, last {bt_summary['Days']} days", sub)
    ws["A4"], ws["B4"] = "Exceptions", int(bt_summary["Exceptions"])
    ws["A5"], ws["B5"] = "Expected", bt_summary["Expected"]
    ws["A6"], ws["B6"] = "Basel zone", bt_summary["Basel zone"]
    ws["B6"].fill = ZONE_FILL[bt_summary["Basel zone"]]
    ws["A7"], ws["B7"] = "Kupiec p-value", bt_summary["Kupiec p-value"]
    ws["A8"], ws["B8"] = "Christoffersen p-value", bt_summary["Indep. p-value"]
    for row in range(4, 9):
        ws.cell(row=row, column=1).font = BOLD
    _header(ws, 10, ["Exception date", "Actual P&L", "VaR forecast", "Excess loss"])
    exc = bt_daily.query("exception")
    for i, (d, row) in enumerate(exc.iterrows(), start=11):
        ws.cell(row=i, column=1, value=d.date()).number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=2, value=float(row["pnl"])).number_format = INR
        ws.cell(row=i, column=3, value=float(row["var"])).number_format = INR
        ws.cell(row=i, column=4, value=f"=-B{i}-C{i}").number_format = INR
    _widths(ws, [24, 14, 14, 14])

    for sheet in wb:
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.sheet_view.showGridLines = False
    wb.save(path)
